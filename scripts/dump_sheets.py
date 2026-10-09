#!/usr/bin/env python3
"""
dump_sheets.py —— xlsx/xlsm/csv 全表转 CSV（openpyxl 兜底通道）

用途：当 editor_sdk MCP 不可用时，用它把工作簿的每个 sheet 导出为 CSV，
      供 personal-financial-analysis-skill 解析。

特性：
  - data_only=True 读公式缓存值（Excel 保存过的文件才有缓存）
  - 自动裁剪尾部全空行/列
  - 保留 #DIV/0! 等错误值文本（它们本身是诊断信号，不能丢）

用法：
  python3 dump_sheets.py <文件路径>                 # 打印全部 sheet
  python3 dump_sheets.py <文件路径> --sheet 0        # 只打印第 0 个 sheet
  python3 dump_sheets.py <文件路径> --list           # 只列 sheet 名与尺寸
  python3 dump_sheets.py <文件路径> --grep 总计      # 只打印含关键词的行（提取趋势用）
  python3 dump_sheets.py <文件路径> --out /tmp/dump  # 每个 sheet 写一个 csv 文件

依赖：openpyxl。缺失时安装到隔离 venv：
  PY=/Users/lyman/.workbuddy/binaries/python/versions/3.13.12/bin/python3
  $PY -m venv /Users/lyman/.workbuddy/binaries/python/envs/default
  /Users/lyman/.workbuddy/binaries/python/envs/default/bin/pip install openpyxl
"""
import sys
import os
import csv
import io
import argparse


def load(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.csv', '.tsv'):
        delim = '\t' if ext == '.tsv' else ','
        with io.open(path, encoding='utf-8-sig', newline='') as f:
            rows = [r for r in csv.reader(f, delimiter=delim)]
        return [('sheet1', rows)]

    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit('缺少 openpyxl。见本脚本 docstring 的隔离 venv 安装方式。')

    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(['' if c is None else str(c) for c in row])
        out.append((ws.title, trim(rows)))
    wb.close()
    return out


def trim(rows):
    """裁掉尾部全空行与全空列"""
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()
    if not rows:
        return rows
    width = max(len(r) for r in rows)
    rows = [r + [''] * (width - len(r)) for r in rows]
    last = -1
    for j in range(width):
        if any(r[j].strip() for r in rows):
            last = j
    return [r[:last + 1] for r in rows]


def to_csv(rows):
    buf = io.StringIO()
    csv.writer(buf, lineterminator='\n').writerows(rows)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('path')
    ap.add_argument('--sheet', type=int, default=None, help='只处理第 N 个 sheet（0-based）')
    ap.add_argument('--list', action='store_true', help='只列 sheet 名与尺寸')
    ap.add_argument('--grep', default=None, help='只打印包含该关键词的行')
    ap.add_argument('--out', default=None, help='输出目录；每 sheet 一个 csv')
    a = ap.parse_args()

    if not os.path.exists(a.path):
        sys.exit('文件不存在：' + a.path)

    sheets = load(a.path)

    if a.list:
        for i, (name, rows) in enumerate(sheets):
            cols = len(rows[0]) if rows else 0
            print(f'[{i}] {name}\trows={len(rows)}\tcols={cols}')
        return

    if a.sheet is not None:
        if not 0 <= a.sheet < len(sheets):
            sys.exit(f'--sheet 越界，共 {len(sheets)} 个 sheet')
        sheets = [sheets[a.sheet]]

    if a.out:
        os.makedirs(a.out, exist_ok=True)
        for i, (name, rows) in enumerate(sheets):
            safe = ''.join(ch if ch.isalnum() or ch in '-_' else '_' for ch in name)
            fp = os.path.join(a.out, f'{i:02d}_{safe}.csv')
            io.open(fp, 'w', encoding='utf-8').write(to_csv(rows))
            print('写入', fp)
        return

    for i, (name, rows) in enumerate(sheets):
        print(f'=== [{i}] {name} ===')
        if a.grep:
            for r in rows:
                line = ','.join(r)
                if a.grep in line:
                    print(line)
        else:
            print(to_csv(rows), end='')
        print()


if __name__ == '__main__':
    main()
